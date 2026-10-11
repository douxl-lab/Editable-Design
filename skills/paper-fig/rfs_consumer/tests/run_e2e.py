"""Real host-runtime/PaperBanana export, native audit, edit, reopen and render.

No mocks. Requires actual artifact-tool and bundled renderer paths. All data is
synthetic; the 1px transport reference cannot establish visual fidelity.
"""
from __future__ import annotations
import argparse
import copy
import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

HERE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(HERE))
import consumer as c
from fixtures import spec,bundle,plan
from protocol.interop import descriptor
from ooxml import NS


def command(args, log, cwd=None):
    result=subprocess.run(args,capture_output=True,text=True,timeout=180,env={**os.environ,'RUNTIME_PYTHON':sys.executable},cwd=cwd)
    Path(log).write_text(result.stdout+'\n'+result.stderr,encoding='utf-8')
    if result.returncode:
        raise RuntimeError('Command failed; see '+str(log)+': '+result.stderr[-2500:])
    return result


def mutate_pptx(source,target,part,fn):
    with zipfile.ZipFile(source) as z:
        payloads=[(i,z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(target,'w') as z:
        for i,raw in payloads:
            if i.filename==part: raw=fn(raw)
            z.writestr(i,raw)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',required=True,type=Path)
    parser.add_argument('--paperbanana',required=True,type=Path)
    parser.add_argument('--research-figure',type=Path)
    args=parser.parse_args()
    output=args.output.resolve(); output.mkdir(parents=True,exist_ok=False)
    results={'synthetic_only':True,'network_calls':0,'paid_generation_calls':0,'visual_fidelity':'not claimed for 1px test references','stages':{},'negative_probes':{}}
    c.write_json(output/'protocol.json',descriptor())
    # Pin protocol files, not unrelated repository documents or version dates.
    for relative in ['contract.py','interop.py','schemas/figure-spec-v1.schema.json','schemas/manifest-v1.schema.json']:
        actual=(args.paperbanana/'rfs'/relative).read_bytes()
        if actual!=(HERE/'protocol'/relative).read_bytes(): raise RuntimeError('Paired PaperBanana protocol drift: '+relative)
        if args.research_figure and actual!=(args.research_figure/'research-figure/protocol'/relative).read_bytes(): raise RuntimeError('RFS protocol drift: '+relative)
    results['stages']['paired_protocol_bytes']='passed'
    command([os.environ.get('RUNTIME_NODE','node'),str(HERE/'artifact.mjs'),'doctor',str(output/'doctor.json')],output/'doctor.log')
    for name,kind in [('branch-feedback-formula','diagram'),('native-bar','plot'),('native-line','plot'),('native-string-categories','plot'),('native-negative-bar','plot'),('native-mixed-bar','plot'),('explicit-line-axis','plot')]:
        root=output/name; root.mkdir()
        scientific=spec(kind)
        if name in {'native-line','explicit-line-axis'}: scientific['datasets'][0]['encoding']['chart_type']='line'
        if name=='native-string-categories':
            scientific['datasets'][0]['rows']=[['1',.3],['01',.7],['=1+1',.5]]
            table={key:scientific['datasets'][0][key] for key in ('columns','rows')}
            content=c.canonical(table); scientific['datasets'][0]['sha256']=c.digest(content.encode())
            scientific['sources'][-1].update(content=content,sha256=c.digest(content.encode()))
        if name in {'native-negative-bar','native-mixed-bar'}:
            rows=scientific['datasets'][0]['rows']
            for i,row in enumerate(rows):
                if name=='native-negative-bar' or i==0: row[1]=-row[1]
            table={key:scientific['datasets'][0][key] for key in ('columns','rows')}
            content=c.canonical(table); scientific['datasets'][0]['sha256']=c.digest(content.encode())
            scientific['sources'][-1].update(content=content,sha256=c.digest(content.encode()))
        manifest=bundle(root/'producer',scientific,args.paperbanana)
        authored=plan(scientific,manifest)
        if name=='explicit-line-axis':
            next(o for o in authored['objects'] if o['role']=='chart')['y_axis_range']=[.2,.8]
        c.write_json(root/'spec.input.json',scientific); c.write_json(root/'plan.json',authored)
        c.prepare(root/'producer',root/'plan.json',root/'consumer',root/'spec.input.json',output/'protocol.json')
        pptx=root/'figure.pptx'
        command([os.environ.get('RUNTIME_NODE','node'),str(HERE/'artifact.mjs'),'build','consumer','figure.pptx'],root/'build.log',cwd=root)
        native=c.audit_pptx(pptx,root/'consumer/reconstruction.json'); c.write_json(root/'native-audit.json',native)
        command([sys.executable,str(HERE/'consumer.py'),'render',str(pptx),str(root/'consumer/reconstruction.json'),str(root/'render')],root/'render.log')
        results['stages'][name]={'native_objects':native['object_count'],'real_pptx_export':True,'reopened':True,'final_file_rendered':True,'relative_cli_paths':True}
        if kind=='diagram':
            repair={'scope':'local','reason':'Synthetic bounded editable-label check','reviewer':'runtime integration test',
                    'before_plan_sha256':c.digest(c.canonical(authored).encode()),'changes':[{'object_name':next(o['name'] for o in authored['objects'] if o['role']=='formula'),'text_color':'#2F6B9A','font_size_pt':13}]}
            c.write_json(root/'repair.json',repair)
            command([os.environ.get('RUNTIME_NODE','node'),str(HERE/'artifact.mjs'),'edit','consumer','figure.pptx','repair.json','edited.pptx'],root/'edit.log',cwd=root)
            c.audit_pptx(root/'edited.pptx',root/'consumer/edited-reconstruction.json')
            command([sys.executable,str(HERE/'consumer.py'),'render',str(root/'edited.pptx'),str(root/'consumer/edited-reconstruction.json'),str(root/'edited-render')],root/'edited-render.log')
            results['stages']['local-edit']={'scientific_fingerprint_unchanged':True,'reopened_after_save':True,'final_file_rendered':True,'relative_cli_paths':True}
    if args.research_figure:
        root=output/'native-handoff'; root.mkdir()
        original=output/'branch-feedback-formula'
        notes=root/'observed-notes.md'
        notes.write_text("Inspected synthetic reference pixels: composition is a one-pixel white transport placeholder. "
                         "No text objects, connector arrows or ports are visible; the palette is white only. "
                         "The native layout is an explicit synthetic test change, with no visual fidelity acceptance.")
        # Fresh process imports the producer's full protocol, not our vendored minimal contract module.
        create="import sys; sys.path.insert(0,sys.argv[1]); from bridge import prepare_editable; prepare_editable(*sys.argv[2:],backend='paper-fig',approved=True)"
        command([sys.executable,'-c',create,str(args.research_figure.resolve()/'research-figure/scripts'),
                 str(original/'spec.input.json'),str(original/'producer'),str(notes),str(root/'handoff')],root/'rfs-handoff.log')
        raw_manifest=(root/'handoff/manifest.json').read_bytes(); manifest=c.strict_json(raw_manifest)
        authored=c.strict_json((original/'plan.json').read_bytes())
        notes_hash=manifest['files_sha256']['observed-design-notes.md']
        authored['reference_review'].update(design_sha256=notes_hash,observed_design_notes_sha256=notes_hash)
        c.write_json(root/'plan.json',authored)
        command([sys.executable,str(HERE/'consumer.py'),'prepare-handoff','handoff','plan.json','consumer',
                 '--source-bundle',str(original/'producer'),'--expected-spec',str(original/'spec.input.json'),
                 '--protocol',str(output/'protocol.json')],root/'prepare.log',cwd=root)
        command([os.environ.get('RUNTIME_NODE','node'),str(HERE/'artifact.mjs'),'build','consumer','figure.pptx'],root/'build.log',cwd=root)
        native=c.audit_pptx(root/'figure.pptx',root/'consumer/reconstruction.json'); c.write_json(root/'native-audit.json',native)
        command([sys.executable,str(HERE/'consumer.py'),'render',str(root/'figure.pptx'),str(root/'consumer/reconstruction.json'),str(root/'render')],root/'render.log')
        if (root/'handoff/manifest.json').read_bytes()!=raw_manifest or (root/'consumer/native-handoff-snapshot/manifest.json').read_bytes()!=raw_manifest:
            raise AssertionError('Original native handoff manifest changed')
        receipt=c.strict_json((root/'consumer/consumer-receipt.json').read_bytes())
        if manifest['status']!='READY_NOT_EXECUTED' or manifest['native_file_verified'] is not False or manifest['scientific_review']!='PENDING':
            raise AssertionError('Producer falsely changed native or scientific status')
        repair={'scope':'local','reason':'Synthetic native-handoff editable-label check','reviewer':'runtime integration test',
                'before_plan_sha256':c.digest(c.canonical(authored).encode()),
                'changes':[{'object_name':next(o['name'] for o in authored['objects'] if o['role']=='formula'),'font_size_pt':13}]}
        c.write_json(root/'repair.json',repair)
        command([os.environ.get('RUNTIME_NODE','node'),str(HERE/'artifact.mjs'),'edit','consumer','figure.pptx','repair.json','edited.pptx'],root/'edit.log',cwd=root)
        c.audit_pptx(root/'edited.pptx',root/'consumer/edited-reconstruction.json')
        command([sys.executable,str(HERE/'consumer.py'),'render',str(root/'edited.pptx'),str(root/'consumer/edited-reconstruction.json'),str(root/'edited-render')],root/'edited-render.log')
        results['stages']['native-handoff']={'actual_rfs_prepare_editable':True,'source_format':receipt['source_format'],
              'source_status_preserved':manifest['status'],'source_native_verified_preserved':manifest['native_file_verified'],
              'source_scientific_review_preserved':manifest['scientific_review'],'approval_bound':receipt['approval_binding'],
              'native_objects':native['object_count'],'reopened':True,'final_file_rendered':True,'native_local_edit_reopened_rendered':True,'visual_review':'pending'}
    probes=output/'negative-probes'; probes.mkdir()
    diagram=output/'branch-feedback-formula'; plot=output/'native-bar'
    def xml(fn):
        def apply(raw):
            root=ET.fromstring(raw); fn(root); return ET.tostring(root)
        return apply
    def probe(name,source,reconstruction,part,fn):
        p=probes/(name+'.pptx'); mutate_pptx(source,p,part,fn)
        try: c.audit_pptx(p,reconstruction)
        except c.Rejected as error: results['negative_probes'][name]={'rejected':True,'error':str(error)}
        else: raise AssertionError('Accepted actual native tampering: '+name)
        return p
    recon=diagram/'consumer/reconstruction.json'; slide='ppt/slides/slide1.xml'
    def change_formula(root):
        for shape in root.findall('.//p:sp',NS):
            nv=shape.find('p:nvSpPr/p:cNvPr',NS)
            if nv is not None and nv.get('name','').endswith(':formula'):
                shape.find('.//a:t',NS).text='invented formula'; return
        raise AssertionError('No native formula was exported')
    probe('changed-formula',diagram/'figure.pptx',recon,slide,xml(change_formula))
    for hidden in ('true','1'):
        def hide_formula(root, hidden=hidden):
            for shape in root.findall('.//p:sp',NS):
                nv=shape.find('p:nvSpPr/p:cNvPr',NS)
                if nv is not None and nv.get('name','').endswith(':formula'):
                    nv.set('hidden',hidden); return
            raise AssertionError('No native formula was exported')
        probe('hidden-formula-'+hidden,diagram/'figure.pptx',recon,slide,xml(hide_formula))
    probe('changed-target-port',diagram/'figure.pptx',recon,slide,xml(lambda root: root.find('.//a:endCxn',NS).set('idx','0')))
    probe('removed-target-arrow',diagram/'figure.pptx',recon,slide,xml(lambda root: root.find('.//a:tailEnd',NS).set('type','none')))
    probe('unknown-object',diagram/'figure.pptx',recon,slide,xml(lambda root: root.find('.//p:spTree',NS).append(copy.deepcopy(root.find('.//p:sp',NS)))))
    def change_order(root):
        tree=root.find('.//p:spTree',NS); item=tree[2]; tree.remove(item); tree.append(item)
    probe('changed-z-order',diagram/'figure.pptx',recon,slide,xml(change_order))
    bad=probe('changed-source-endpoint',diagram/'figure.pptx',recon,slide,xml(lambda root: root.find('.//a:stCxn',NS).set('id','999')))
    # Editing refuses a mutated imported PPTX before authoring another output.
    failed=subprocess.run([os.environ.get('RUNTIME_NODE','node'),str(HERE/'artifact.mjs'),'edit',str(diagram/'consumer'),str(bad),str(diagram/'repair.json'),str(probes/'must-not-exist.pptx')],capture_output=True,text=True,env={**os.environ,'RUNTIME_PYTHON':sys.executable})
    if failed.returncode==0 or (probes/'must-not-exist.pptx').exists(): raise AssertionError('Tampered edit input was accepted')
    results['negative_probes']['tampered-edit-input']={'rejected':True}
    with zipfile.ZipFile(plot/'figure.pptx') as z:
        chartpart=next(n for n in z.namelist() if re_chart(n))
        workbook=next(n for n in z.namelist() if n.endswith('.xlsx'))
    def truncate_axis(root):
        minimum=root.find('.//c:valAx/c:scaling/c:min',NS)
        minimum.set('val','0.2')
    probe('truncated-bar-baseline',plot/'figure.pptx',plot/'consumer/reconstruction.json',chartpart,xml(truncate_axis))
    probe('changed-chart-cache' ,plot/'figure.pptx',plot/'consumer/reconstruction.json',chartpart,xml(lambda root: setattr(root.find('.//c:val//c:v',NS),'text','99')))
    def workbook_mutation(raw):
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(raw)); book.worksheets[0]['B2']=99; out=io.BytesIO(); book.save(out); return out.getvalue()
    probe('changed-chart-workbook',plot/'figure.pptx',plot/'consumer/reconstruction.json',workbook,workbook_mutation)
    line=output/'native-line'
    with zipfile.ZipFile(line/'figure.pptx') as z:
        line_part=next(n for n in z.namelist() if re_chart(n))
    def smooth_line(root):
        plot=root.find('.//c:lineChart',NS)
        smooth=plot.find('c:smooth',NS)
        if smooth is None: smooth=ET.SubElement(plot,'{'+NS['c']+'}smooth')
        smooth.set('val','1')
    probe('unrequested-smoothed-line',line/'figure.pptx',line/'consumer/reconstruction.json',line_part,xml(smooth_line))
    def changed_series_color(root):
        root.find('.//c:ser/c:spPr/a:ln/a:solidFill/a:srgbClr',NS).set('val','FF0000')
    probe('changed-line-series-color',line/'figure.pptx',line/'consumer/reconstruction.json',line_part,xml(changed_series_color))
    categorical=output/'native-string-categories'
    with zipfile.ZipFile(categorical/'figure.pptx') as z:
        categorical_workbook=next(n for n in z.namelist() if n.endswith('.xlsx'))
    def category_type_mutation(raw):
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(raw)); book.worksheets[0]['A2']=1
        out=io.BytesIO(); book.save(out); return out.getvalue()
    probe('numeric-coercion-of-string-category',categorical/'figure.pptx',categorical/'consumer/reconstruction.json',categorical_workbook,category_type_mutation)
    def zero_loss_mutation(raw):
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(raw)); book.worksheets[0]['A3']='1'
        out=io.BytesIO(); book.save(out); return out.getvalue()
    probe('lost-leading-zero-category',categorical/'figure.pptx',categorical/'consumer/reconstruction.json',categorical_workbook,zero_loss_mutation)
    c.write_json(output/'results.json',results)
    print(json.dumps(results,indent=2))


def re_chart(n):
    return '/charts/chart' in n and n.endswith('.xml')


if __name__=='__main__': main()
